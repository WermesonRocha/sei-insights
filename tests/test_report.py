import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from sei_insights.report import build_resumo, write_spreadsheet
from sei_insights.store import ProcessRow


def row(numero: str, situacao: str = "Na CTI", status: str = "concluído") -> ProcessRow:
    return ProcessRow(
        numero=numero, titulo="t", data_execucao="2026-09-22 10:00:00",
        data_analise="2026-09-22 10:00:00", data_ultimo_despacho="",
        situacao=situacao, destino="", acao_esperada="", pendencia_curta="",
        link_process="url", status_coleta=status, hash_ultimo_despacho="h",
    )


class ReportTest(unittest.TestCase):
    def test_write_leitura_de_volta(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.xlsx"
            rows = [row("001", "Na CTI"), row("002", "Pendente de assinatura")]
            novos = [row("002", "Pendente de assinatura")]
            resumo = build_resumo(rows, novos)
            write_spreadsheet(path, rows, novos, resumo)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Aba principal", "Novos", "Resumo"])
            principal = wb["Aba principal"]
            self.assertEqual(principal.max_row, len(rows) + 1)
            self.assertEqual(principal.cell(row=1, column=1).value, "numero")
            self.assertEqual(principal.cell(row=2, column=1).value, "001")
            nov = wb["Novos"]
            self.assertEqual(nov.max_row, len(novos) + 1)
            res = wb["Resumo"]
            self.assertEqual(res["B2"].value, 2)   # total
            self.assertEqual(res["B3"].value, 1)   # novos
            wb.close()

    def test_resultado_vazio_escreve_workbook_valido(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "vazio.xlsx"
            write_spreadsheet(path, [], [], build_resumo([], []))
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Aba principal", "Novos", "Resumo"])
            self.assertEqual(wb["Aba principal"].max_row, 1)
            self.assertEqual(wb["Resumo"]["B2"].value, 0)
            wb.close()