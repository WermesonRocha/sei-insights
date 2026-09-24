from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from sei_insights.storage.mirror import FIELDS, ProcessRow


def build_resumo(rows: list[ProcessRow], novos: list[ProcessRow]) -> dict:
    por_situacao: dict[str, int] = {}
    por_status: dict[str, int] = {}
    for r in rows:
        por_situacao[r.situacao] = por_situacao.get(r.situacao, 0) + 1
        por_status[r.status_coleta] = por_status.get(r.status_coleta, 0) + 1
    return {
        "total": len(rows),
        "novos": len(novos),
        "por_situacao": por_situacao,
        "por_status": por_status,
    }


def _fill_sheet(sheet, rows: list[ProcessRow]) -> None:
    sheet.append(FIELDS)
    for r in rows:
        sheet.append([getattr(r, f) for f in FIELDS])


def write_spreadsheet(
    path: Path,
    rows: list[ProcessRow],
    novos: list[ProcessRow],
    resumo: dict,
) -> None:
    wb = Workbook()
    principal = wb.active
    principal.title = "Aba principal"
    _fill_sheet(principal, rows)
    nov = wb.create_sheet("Novos")
    _fill_sheet(nov, novos)
    res = wb.create_sheet("Resumo")
    res.append(["métrica", "valor"])
    res.append(["total", resumo["total"]])
    res.append(["novos", resumo["novos"]])
    for label, counts in (("por_situacao", resumo["por_situacao"]),
                          ("por_status", resumo["por_status"])):
        for key, value in counts.items():
            res.append([label, f"{key} = {value}"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    wb.close()